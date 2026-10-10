"""
Tripura Job Pulse - daily job finder

Searches news feeds for new Tripura recruitment notices, fills in what it can
reliably work out, and emails you an Excel file in your upload format.

Be clear about what this does and doesn't do. Headlines and short snippets
give a title, an organisation and a link. They do not give vacancy counts,
pay scales or age limits - those sit inside PDF notifications. Those columns
come back blank, shaded yellow, for you to complete from the official notice.

Environment variables:
  SUPABASE_URL
  SUPABASE_SECRET_KEY
  GMAIL_USER
  GMAIL_APP_PASSWORD
  MAIL_TO              where to send it (defaults to GMAIL_USER)
  LOOKBACK_HOURS       how far back to look (default 48)
  DRY_RUN              "true" to print instead of emailing
"""

import html
import os
import re
import smtplib
import ssl
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import formataddr
from io import BytesIO

import feedparser
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from supabase import create_client

try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:
    pass

SUPABASE_URL = os.environ["SUPABASE_URL"].rstrip("/")
SUPABASE_KEY = os.environ["SUPABASE_SECRET_KEY"]
GMAIL_USER = os.environ["GMAIL_USER"]
GMAIL_PASS = os.environ["GMAIL_APP_PASSWORD"]
MAIL_TO = os.environ.get("MAIL_TO") or GMAIL_USER
LOOKBACK_HOURS = int(os.environ.get("LOOKBACK_HOURS", "48"))
DRY_RUN = os.environ.get("DRY_RUN", "false").lower() == "true"

sb = create_client(SUPABASE_URL, SUPABASE_KEY)

COLUMNS = ["Job title", "Organisation", "Job type", "City", "State",
           "Experience", "Industry", "Qualification", "Vacancies", "Salary",
           "Age limit", "Last date", "How to apply", "Apply link", "Description"]

# Columns the script can usually fill. The rest are shaded for you to complete.
AUTO = {"Job title", "Organisation", "Job type", "City", "State",
        "Apply link", "Description"}

# Job board feeds carry actual vacancy posts; news feeds mostly carry stories
# about jobs. Both are kept, but the job boards do the real work.
# Each entry is (label, url). Dead feeds are reported and skipped.
FEEDS = [
    # Tripura job boards - published RSS, meant to be read by other software
    ("tripurajob.in", "https://www.tripurajob.in/feeds/posts/default?alt=rss"),
    ("tripuracareer.in", "https://www.tripuracareer.in/feeds/posts/default?alt=rss"),
    ("jobstripura.com", "https://www.jobstripura.com/feeds/posts/default?alt=rss"),
    ("tripurastarnews.com", "https://www.tripurastarnews.com/category/job-employment/feed/"),

    # News searches - thinner, but occasionally first with a big announcement
    ("news: Tripura recruitment",
     "https://news.google.com/rss/search?q=Tripura+recruitment&hl=en-IN&gl=IN&ceid=IN:en"),
    ("news: Tripura apply online",
     "https://news.google.com/rss/search?q=Tripura+vacancy+apply+online&hl=en-IN&gl=IN&ceid=IN:en"),
    ("news: TPSC",
     "https://news.google.com/rss/search?q=TPSC+recruitment&hl=en-IN&gl=IN&ceid=IN:en"),
    ("news: JRBT",
     "https://news.google.com/rss/search?q=JRBT+recruitment+Tripura&hl=en-IN&gl=IN&ceid=IN:en"),
    ("news: Agartala vacancy",
     "https://news.google.com/rss/search?q=Agartala+job+vacancy&hl=en-IN&gl=IN&ceid=IN:en"),
    ("news: walk-in",
     "https://news.google.com/rss/search?q=Tripura+walk-in+interview&hl=en-IN&gl=IN&ceid=IN:en"),
]

# A headline has to look like a job notice, not general news.
JOB_WORDS = re.compile(
    r"\b(recruit\w*|vacanc\w*|vacancies|hiring|appointment|apply\s+online|"
    r"walk[-\s]?in|notification|post[s]?\b|job[s]?\b|engage\w*|"
    r"empanel\w*|advertis\w*)\b", re.I)

# Things that mention jobs but aren't a vacancy to apply for.
SKIP_WORDS = re.compile(
    r"\b(result|answer\s*key|admit\s*card|cut[-\s]?off|merit\s*list|"
    r"exam\s*date|postponed|cancell?ed|scam|fraud|protest|unemployment\s+rate|"
    r"job\s*loss|laid\s*off|retrench\w*)\b", re.I)

# News *about* jobs reads very differently from a notice. A pay-gap story or a
# court hearing mentions vacancies without being one, so those get dropped.
COMMENTARY = re.compile(
    r"\b(high\s+court|supreme\s+court|affidavit|petition|plea|hearing|verdict|"
    r"listed\s+for|seeks?|slams?|alleges?|expose[sd]?|pay\s+gap|row\s+over|"
    r"demands?|criticis\w*|questions?|debate|assembly|minister\s+said|"
    r"opposition|survey|report\s+says|study|analysis|crisis|backlog)\b", re.I)

# A real notice almost always says one of these.
ANNOUNCEMENT = re.compile(
    r"(\brecruitment\b|\bnotification\b|\bapply\s+online\b|\bapply\s+before\b|"
    r"\bwalk[-\s]?in\b|\binvites?\s+application|\bapplications?\s+invited\b|"
    r"\blast\s+date\b|\bvacanc\w*\s+(?:announced|out|notified)\b|"
    r"\b\d{1,5}\+?\s+(?:post|posts|vacanc\w*)\b|"
    r"\brecruitment\s+\d{4}\b|\bhiring\s+for\b)", re.I)

# Either a Tripura place name, or a body that only exists in Tripura - a TPSC
# headline often never says the word "Tripura".
RELEVANT_PLACE = re.compile(
    r"\b(tripura|agartala|udaipur|dharmanagar|kailashahar|belonia|ambassa|"
    r"khowai|sabroom|tpsc|jrbt|trbt|ttaadc|tsecl|agmc|trtc)\b", re.I)

ORGS = [
    ("TPSC", "Tripura Public Service Commission (TPSC)"),
    ("JRBT", "Joint Recruitment Board Tripura (JRBT)"),
    ("TRBT", "Teachers' Recruitment Board Tripura"),
    ("Tripura University", "Tripura University"),
    ("NIT Agartala", "NIT Agartala"),
    ("TSECL", "Tripura State Electricity Corporation (TSECL)"),
    ("TTAADC", "Tripura Tribal Areas Autonomous District Council"),
    ("Tripura Police", "Tripura Police"),
    ("AGMC", "Agartala Government Medical College"),
    ("TIDC", "Tripura Industrial Development Corporation"),
    ("TRTC", "Tripura Road Transport Corporation"),
    ("Tripura High Court", "Tripura High Court"),
]

GOV_HINT = re.compile(
    r"\b(government|govt|department|ministry|commission|board|council|"
    r"corporation|directorate|municipal|nagar|panchayat|district|"
    r"tpsc|jrbt|trbt|police|university|college|hospital|railway|"
    r"bank|ssc|upsc|court)\b", re.I)


def clean(text):
    text = html.unescape(text or "")
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def strip_publisher(title):
    """Google News appends ' - Publisher' to every headline."""
    return re.sub(r"\s+-\s+[^-]{2,40}$", "", title).strip()


def entry_time(entry):
    parsed = entry.get("published_parsed") or entry.get("updated_parsed")
    if parsed:
        return datetime(*parsed[:6], tzinfo=timezone.utc)
    return datetime.now(timezone.utc)


def guess_organisation(text):
    for needle, full in ORGS:
        if re.search(r"\b" + re.escape(needle) + r"\b", text, re.I):
            return full
    return ""


def guess_city(text):
    for city in ["Agartala", "Udaipur", "Dharmanagar", "Kailashahar",
                 "Belonia", "Ambassa", "Khowai", "Sabroom"]:
        if re.search(r"\b" + city + r"\b", text, re.I):
            return city
    return ""


def guess_job_type(text):
    return "Government" if GOV_HINT.search(text) else ""


def known_urls():
    rows = sb.table("job_leads").select("source_url").limit(10000).execute().data or []
    return {r["source_url"] for r in rows}


def collect():
    cutoff = datetime.now(timezone.utc) - timedelta(hours=LOOKBACK_HOURS)
    seen_urls = known_urls()
    seen_titles = set()
    found = []
    tally = []

    for label, url in FEEDS:
        kept_here = 0
        # Why things were dropped, so a feed returning nothing can be explained
        # rather than guessed at.
        why = {"seen": 0, "duplicate": 0, "too old": 0, "not a job": 0,
               "skipped word": 0, "commentary": 0, "no notice wording": 0,
               "not Tripura": 0}
        try:
            feed = feedparser.parse(url)
        except Exception as err:
            tally.append((label, "could not be read", 0, 0, why))
            print(f"  ! {label}: {err}")
            continue

        total_here = len(feed.entries)
        for e in feed.entries:
            link = e.get("link", "")
            raw_title = clean(e.get("title", ""))
            if not link or not raw_title:
                continue
            if link in seen_urls:
                why["seen"] += 1
                continue

            title = strip_publisher(raw_title)
            key = re.sub(r"[^a-z0-9]", "", title.lower())
            if key in seen_titles:
                why["duplicate"] += 1
                continue

            snippet = clean(e.get("summary") or e.get("description") or "")[:800]
            blob = f"{title} {snippet}"

            if not JOB_WORDS.search(blob):
                why["not a job"] += 1
                continue
            if SKIP_WORDS.search(title):
                why["skipped word"] += 1
                continue
            if COMMENTARY.search(title):
                why["commentary"] += 1
                continue
            if not ANNOUNCEMENT.search(blob):
                why["no notice wording"] += 1
                continue
            if not RELEVANT_PLACE.search(blob):
                why["not Tripura"] += 1
                continue
            if entry_time(e) < cutoff:
                why["too old"] += 1
                continue

            src = e.get("source", {})
            found.append({
                "title": title,
                "snippet": snippet,
                "link": link,
                "source_name": (src.get("title") if isinstance(src, dict) else "") or label,
                "when": entry_time(e),
            })
            seen_urls.add(link)
            seen_titles.add(key)
            kept_here += 1

        tally.append((label, "ok" if total_here else "empty",
                      total_here, kept_here, why))

    found.sort(key=lambda x: x["when"], reverse=True)
    return found, tally


def to_row(item):
    blob = f"{item['title']} {item['snippet']}"
    org = guess_organisation(blob)
    # A recognised Tripura body is a government employer, whatever the wording.
    job_type = "Government" if org else guess_job_type(blob)
    return {
        "Job title": item["title"],
        "Organisation": org or (item["source_name"] or ""),
        "Job type": job_type,
        "City": guess_city(blob),
        "State": "Tripura" if RELEVANT_PLACE.search(blob) else "",
        "Experience": "",
        "Industry": "",
        "Qualification": "",
        "Vacancies": "",
        "Salary": "",
        "Age limit": "",
        "Last date": "",
        "How to apply": "",
        "Apply link": item["link"],
        "Description": item["snippet"],
    }


def build_excel(rows, found_on):
    wb = Workbook()
    ws = wb.active
    ws.title = "New jobs"

    head_fill = PatternFill("solid", fgColor="0E4F4C")
    todo_fill = PatternFill("solid", fgColor="FDF3D7")
    head_font = Font(color="FFFFFF", bold=True, size=11)

    ws.append(COLUMNS)
    for i, name in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i)
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in rows:
        ws.append([r[c] for c in COLUMNS])

    # Shade the columns you need to complete by hand
    for i, name in enumerate(COLUMNS, start=1):
        if name in AUTO:
            continue
        for row in range(2, len(rows) + 2):
            ws.cell(row=row, column=i).fill = todo_fill

    widths = {"Job title": 44, "Organisation": 30, "Job type": 13, "City": 14,
              "State": 12, "Experience": 18, "Industry": 18, "Qualification": 26,
              "Vacancies": 11, "Salary": 14, "Age limit": 16, "Last date": 13,
              "How to apply": 24, "Apply link": 40, "Description": 60}
    for i, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 18)

    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    notes = wb.create_sheet("Read me")
    for line in [
        ["Tripura Job Pulse - jobs found on " + found_on],
        [],
        ["White columns were filled automatically from the news item."],
        ["Yellow columns could not be worked out from a headline."],
        ["Open the Apply link, read the official notice, and fill those in."],
        [],
        ["Check every row against the official notification before you upload it."],
        ["Job boards and news sites get details wrong, and some items are not"],
        ["real vacancies."],
        [],
        ["The Apply link points at wherever the notice was found, which is often"],
        ["a job board rather than the official site. Replace it with the official"],
        ["application link before you upload."],
        [],
        ["Delete any row that is not a genuine opening, then upload the file"],
        ["through Admin - Jobs on tripurajobpulse.com."],
    ]:
        notes.append(line)
    notes.column_dimensions["A"].width = 95

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def remember(items):
    rows = [{
        "source_url": i["link"],
        "title": i["title"],
        "organisation": guess_organisation(f"{i['title']} {i['snippet']}"),
        "source_name": i["source_name"],
        "reported": True,
    } for i in items]
    if rows:
        sb.table("job_leads").upsert(
            rows, on_conflict="source_url", ignore_duplicates=True
        ).execute()


def send(xlsx, count, found_on):
    subject = f"{count} possible Tripura job{'' if count == 1 else 's'} - {found_on}"
    body = (
        f"{count} new Tripura job notice{'' if count == 1 else 's'} turned up today.\n\n"
        "The attached file is in your upload format. White columns are filled in; "
        "yellow ones need you to open the link and read the official notice.\n\n"
        "Check each row before uploading - news sites get details wrong, and the "
        "occasional item is not a real vacancy.\n\n"
        "Tripura Job Pulse"
    )

    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = formataddr(("Tripura Job Pulse", GMAIL_USER))
    msg["To"] = MAIL_TO
    msg.set_content(body)
    msg.add_attachment(
        xlsx,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"tripura-jobs-{found_on}.xlsx",
    )

    with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ssl.create_default_context()) as s:
        s.login(GMAIL_USER, GMAIL_PASS)
        s.send_message(msg)


def summary(line):
    print(line)
    path = os.environ.get("GITHUB_STEP_SUMMARY")
    if path:
        with open(path, "a", encoding="utf-8") as f:
            f.write(line + "\n\n")


def main():
    found_on = datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d")
    items, tally = collect()

    summary("## Result")

    # Which sources are pulling their weight, and why the rest dropped out
    lines = ["| Source | In feed | Kept | Why the rest were dropped |",
             "| --- | ---: | ---: | --- |"]
    for label, state, total, kept, why in tally:
        shown = total if state == "ok" else state
        reasons = ", ".join(f"{k} {v}" for k, v in why.items() if v) or "—"
        lines.append(f"| {label} | {shown} | {kept} | {reasons} |")
    summary("\n".join(lines))

    if not items:
        summary("No new Tripura job notices found today. Nothing emailed.")
        return

    rows = [to_row(i) for i in items]
    xlsx = build_excel(rows, found_on)
    summary(f"Found **{len(rows)}** possible job{'' if len(rows) == 1 else 's'}:")
    for r in rows[:25]:
        summary(f"- {r['Job title']}")

    if DRY_RUN:
        summary("**Practice run - nothing emailed.**")
        return

    send(xlsx, len(rows), found_on)
    remember(items)
    summary(f"**Emailed to {MAIL_TO}.**")


if __name__ == "__main__":
    main()
